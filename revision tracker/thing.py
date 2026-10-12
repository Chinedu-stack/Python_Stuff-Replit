"""
Add Biology topics to GCSE_Revision_Tracker.xlsx
- Important (high-mark) topics listed FIRST
- Less important topics after
- Keeps existing Dashboard/Settings/formulas intact
- Expands the Biology table if needed
"""

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

FILE = "revision tracker/GCSE_Revision_Tracker.xlsx"

IMPORTANT = [
    "Movements of molecules",
    "Excretion",
    "Human Reproduction",
    "Organisms in the environment",
    "Co-ordination and response",
    "Transport in Humans",
    "Transport in plants",
    "Inheritance",
    "Investigation skills and planning investigations",
]

OTHER = [
    "Characteristics of living organisms",
    "Variety of living organisms",
    "Levels of organisation",
    "Cell structure",
    "Biological Molecules",
    "Movement of substances",
    "Nutrition in humans",
    "Nutrition in plants",
    "Respiration",
    "Gas exchange in humans",
    "Gas exchange in plants",
    "Reproduction in plants",
    "Feeding relationships",
    "Cycles within ecosystems",
    "Human Influences",
]

TOPICS = IMPORTANT + OTHER

# ---- Style constants (must match the builder) ----
NAVY = "1F3864"
BORDER_CL = "BFD4EA"
THIN = Side(style="thin", color=BORDER_CL)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT   = Alignment(horizontal="left", vertical="center")

wb = load_workbook(FILE)
ws = wb["Biology"]

# Find the Biology table
bio_table = None
for tbl in ws.tables.values():
    bio_table = tbl
    break

if bio_table is None:
    raise SystemExit("Could not find the Biology table.")

ref = bio_table.ref  # e.g. "A3:F18"
left, right = ref.split(":")
start_row = int("".join([c for c in left if c.isdigit()]))
end_row   = int("".join([c for c in right if c.isdigit()]))

FIRST_DATA_ROW = start_row + 1
current_rows = end_row - start_row
needed_rows  = len(TOPICS)

if needed_rows > current_rows:
    extra = needed_rows - current_rows
    new_end = end_row + extra

    for r in range(end_row + 1, new_end + 1):
        ws.row_dimensions[r].height = 24
        for c in range(1, 7):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.alignment = LEFT if c in (1, 6) else CENTER
        # Next Review Date formula
        ws.cell(row=r, column=5, value=(
            f'=IF(OR(D{r}="",C{r}="",C{r}="Not Assessed"),"",'
            f'D{r}+IF(C{r}="Red",RedInterval,'
            f'IF(C{r}="Amber",AmberInterval,GreenInterval)))'
        ))
        ws.cell(row=r, column=4).number_format = "DD/MM/YYYY"
        ws.cell(row=r, column=5).number_format = "DD/MM/YYYY"

    bio_table.ref = f"A{start_row}:F{new_end}"
    end_row = new_end

# Write topic names
for i, topic in enumerate(TOPICS):
    ws.cell(row=FIRST_DATA_ROW + i, column=1, value=topic)

wb.save(FILE)
print(f"Added {len(TOPICS)} Biology topics.")
print(f"First {len(IMPORTANT)} are the high-mark topics.")
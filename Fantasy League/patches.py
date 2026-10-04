# sort_market.py  (v3 — prices live, order static)
# Names (B) are static — they define the sort order.
# Prices (C), Change (D), Ownership (E) are live formulas that update when SETTINGS changes.

import os
import sys
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.formatting.rule import FormulaRule

WB_PATH = r"C:\Users\chine\OneDrive\Copies School Fantasy League for editing\11O Fantasy League.xlsx"
SET_PLAYER_START = 22
SET_PLAYER_END   = 51
CALC_START = 2
CALC_END   = 31
PM_START   = 5
PM_END     = 34

WHITE="FFFFFF"; PALE_BLUE="F5F9FF"; DARK="172033"; GREY="6B7280"
LIGHT_GREEN="DCFCE7"; GREEN="16A34A"; LIGHT_RED="FEE2E2"; RED="DC2626"
LIGHT_GREY="E5E7EB"
thin_grey = Side(style="thin", color=LIGHT_GREY)

def box(cell, fill=WHITE, bold=False, color=DARK, align="left"):
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(bold=bold, color=color, size=12)
    cell.alignment = Alignment(horizontal=align, vertical="center")
    cell.border = Border(left=thin_grey, right=thin_grey, top=thin_grey, bottom=thin_grey)

def to_price(v, fallback=0.0):
    if v is None: return fallback
    if isinstance(v, (int, float)): return float(v)
    try:
        return float(str(v).replace("£", "").replace("m", "").strip())
    except (ValueError, TypeError):
        return fallback

if not os.path.exists(WB_PATH):
    print(f"ERROR: {WB_PATH} not found"); sys.exit(1)

print(f"Opening {WB_PATH}")
wb = load_workbook(WB_PATH)
s = wb["SETTINGS"]
m = wb["PLAYER MARKET"]

# Read players + current prices to determine sort order
players = []
for i in range(SET_PLAYER_END - SET_PLAYER_START + 1):
    r = SET_PLAYER_START + i
    name = s.cell(r, 2).value
    if name is None or str(name).strip() == "":
        continue
    base  = to_price(s.cell(r, 3).value, 0.0)
    final = to_price(s.cell(r, 6).value, base)
    players.append({"name": str(name).strip(), "final": final, "order": i})

print(f"Found {len(players)} players — sorting for row order")
players.sort(key=lambda p: (-p["final"], p["order"]))

# Hide helper columns (we don't use them any more, but hide for cleanliness)
for col in ("X", "Y", "Z"):
    m.column_dimensions[col].hidden = True

# Wipe old helpers
for r in range(PM_START, PM_END + 1):
    for c in (24, 25, 26):
        m.cell(r, c).value = None

for i in range(PM_END - PM_START + 1):
    r = PM_START + i

    if i < len(players):
        # B — player NAME (static, defines the row's identity / sort position)
        m.cell(r, 2).value = players[i]["name"]
    else:
        m.cell(r, 2).value = None

    # C — PRICE (live formula from CALCULATIONS)
    m.cell(r, 3).value = (
        f'=IF(B{r}="","",IFERROR(INDEX(CALCULATIONS!$C${CALC_START}:$C${CALC_END},'
        f'MATCH(B{r},CALCULATIONS!$A${CALC_START}:$A${CALC_END},0)),""))'
    )
    m.cell(r, 3).number_format = '£0.0"m"'

    # D — CHANGE (live)
    m.cell(r, 4).value = (
        f'=IF(B{r}="","",IF(CurrentGW<SETTINGS!$C$16,0,'
        f'INDEX(CALCULATIONS!$C${CALC_START}:$C${CALC_END},'
        f'MATCH(B{r},CALCULATIONS!$A${CALC_START}:$A${CALC_END},0))'
        f'-'
        f'INDEX(CALCULATIONS!$B${CALC_START}:$B${CALC_END},'
        f'MATCH(B{r},CALCULATIONS!$A${CALC_START}:$A${CALC_END},0))))'
    )
    m.cell(r, 4).number_format = '"▲ +£"0.0"m";"▼ -£"0.0"m";"— £0.0m"'

    # E — OWNERSHIP (live)
    m.cell(r, 5).value = (
        f'=IF(B{r}="","",COUNTIF(CALCULATIONS!$L${CALC_START}:$P${CALC_END},B{r})'
        f'/COUNTA(CALCULATIONS!$K${CALC_START}:$K${CALC_END}))'
    )
    m.cell(r, 5).number_format = '0%'

    # Styling
    fill = WHITE if i % 2 == 0 else PALE_BLUE
    box(m.cell(r, 2), fill, False, DARK, "left")
    box(m.cell(r, 3), fill, True, DARK, "center")
    box(m.cell(r, 4), fill, True, DARK, "center")
    box(m.cell(r, 5), fill, False, DARK, "center")
    m.row_dimensions[r].height = 28

m.column_dimensions["B"].width = 24
m.column_dimensions["C"].width = 12
m.column_dimensions["D"].width = 14
m.column_dimensions["E"].width = 14

# CF on Change column
keys_to_drop = [k for k in m.conditional_formatting._cf_rules.keys() if str(k).startswith("D")]
for k in keys_to_drop:
    m.conditional_formatting._cf_rules.pop(k, None)

m.conditional_formatting.add(f"D{PM_START}:D{PM_END}", FormulaRule(
    formula=[f'D{PM_START}>0'], fill=PatternFill("solid", fgColor=LIGHT_GREEN),
    font=Font(color=GREEN, bold=True)))
m.conditional_formatting.add(f"D{PM_START}:D{PM_END}", FormulaRule(
    formula=[f'D{PM_START}<0'], fill=PatternFill("solid", fgColor=LIGHT_RED),
    font=Font(color=RED, bold=True)))
m.conditional_formatting.add(f"D{PM_START}:D{PM_END}", FormulaRule(
    formula=[f'AND(D{PM_START}=0,B{PM_START}<>"")'], fill=PatternFill("solid", fgColor="F3F4F6"),
    font=Font(color=GREY, bold=True)))

wb.save(WB_PATH)
print(f"\nSaved: {WB_PATH}")
print(f"PLAYER MARKET has {len(players)} players. Prices are now LIVE.")
print("Order is fixed until you re-run this script.")
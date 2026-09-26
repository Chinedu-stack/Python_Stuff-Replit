from openpyxl import load_workbook

# =========================================================
# CONFIG
# =========================================================
FILE = r"C:\Users\chine\Downloads\Year_11_Fantasy_League (1).xlsx"

wb = load_workbook(FILE)
calc = wb["CALCULATIONS"]

print(f"Loaded: {FILE}")
print("Patching CALCULATIONS column D formulas...")
print()

# Player rows on CALCULATIONS are 2..22.
# Student rows on TEST RESULTS are 5..25.
# Offset = +3 rows.
for r in range(2, 23):
    test_row = r + 3
    formula = (
        f'=_xlfn.IFERROR(IF(\'TEST RESULTS\'!C{test_row}="",0,'
        f'IF(\'TEST RESULTS\'!C{test_row}>=95,10,'
        f'IF(\'TEST RESULTS\'!C{test_row}>=90,9,'
        f'IF(\'TEST RESULTS\'!C{test_row}>=85,8,'
        f'IF(\'TEST RESULTS\'!C{test_row}>=80,7,'
        f'IF(\'TEST RESULTS\'!C{test_row}>=75,6,'
        f'IF(\'TEST RESULTS\'!C{test_row}>=70,5,'
        f'IF(\'TEST RESULTS\'!C{test_row}>=65,4,'
        f'IF(\'TEST RESULTS\'!C{test_row}>=60,3,0))))))))),0)'
    )
    player_name = calc.cell(r, 1).value or "?"
    calc.cell(r, 4).value = formula
    print(f"  Row {r} ({player_name}) → TEST RESULTS!C{test_row}")

wb.save(FILE)
print()
print(f"Done. Saved: {FILE}")
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

# =========================================================
# DATA
# =========================================================
players = sorted([
    ("Lia Li", 13.0),
    ("Chinedu Chile", 11.0),
    ("Emily Warburton", 11.0),
    ("Sophie Barker", 13.0),
    ("Alyza Rai", 10.5),
    ("David Wu", 10.5),
    ("Joe Paige", 10.5),
    ("Azlan Husain", 10.0),
    ("Nina De Saram", 10.0),
    ("Alexa Amesimeku", 9.5),
    ("Daniel Liu", 9.5),
    ("Rebecca Seymour", 9.5),
    ("Edwin Mayers", 9.0),
    ("Melody Zhaeentan", 8.5),
    ("Esme Beston", 8.0),
    ("Iona Smith", 8.0),
    ("Oscar Day-Duran", 8.0),
    ("Alfie Foley", 7.5),
    ("Ava Swinfen", 7.0),
    ("Rory Bowman", 7.0),
    ("Xander Meijer", 7.0),
], key=lambda p: -p[1])
managers = sorted([p[0] for p in players])

# =========================================================
# THEME
# =========================================================
NAVY = "0B1F3A"
BLUE = "1565C0"
LIGHT_BLUE = "EAF2FF"
PALE_BLUE = "F5F9FF"
WHITE = "FFFFFF"
DARK = "172033"
GREY = "6B7280"
LIGHT_GREY = "E5E7EB"
GREEN = "16A34A"
LIGHT_GREEN = "DCFCE7"
RED = "DC2626"
LIGHT_RED = "FEE2E2"
GOLD = "F59E0B"
LIGHT_GOLD = "FEF3C7"
SLATE = "475569"
SLATE_DARK = "334155"

thin_grey = Side(style="thin", color=LIGHT_GREY)
medium_blue = Side(style="medium", color=BLUE)

def title(ws, text, row=1, start_col=1, end_col=8):
    ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
    c = ws.cell(row, start_col, text)
    c.font = Font(size=22, bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 52

def section(ws, text, row, start_col=1, end_col=8, fill=BLUE):
    ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
    c = ws.cell(row, start_col, text)
    c.font = Font(size=13, bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=fill)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 32

def header_row(ws, row, start_col, end_col, fill=NAVY):
    for col in range(start_col, end_col + 1):
        c = ws.cell(row, col)
        c.font = Font(bold=True, color=WHITE, size=12)
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = Border(bottom=medium_blue)
    ws.row_dimensions[row].height = 32

def box(cell, fill=WHITE, bold=False, color=DARK, align="left", size=12, italic=False):
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(bold=bold, color=color, size=size, italic=italic)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=False)
    cell.border = Border(left=thin_grey, right=thin_grey, top=thin_grey, bottom=thin_grey)

wb = Workbook()
wb.remove(wb.active)

# ---------------------------------------------------------
# CALCULATIONS
# ---------------------------------------------------------
calc = wb.create_sheet("CALCULATIONS")
calc.sheet_state = "hidden"

calc["A1"] = "Player"
calc["B1"] = "Base Price"
calc["C1"] = "Current Price"
calc["D1"] = "Points This Week"
calc["E1"] = "Last Week Pts"
calc["F1"] = "Δ"
calc["W1"] = "Dropdown"

for i, (name, price) in enumerate(players, start=2):
    calc.cell(i, 1, name)
    calc.cell(i, 2, price)
    calc.cell(i, 3, f"=MIN(13,MAX(7,B{i}))")
    
    # FIX: Use INDEX/MATCH to find the player by name in TEST RESULTS
    calc.cell(i, 4,
        f'=_xlfn.IFERROR(IF(INDEX(\'TEST RESULTS\'!$C$5:$C$25,MATCH(A{i},\'TEST RESULTS\'!$B$5:$B$25,0))="",0,'
        f'IF(INDEX(\'TEST RESULTS\'!$C$5:$C$25,MATCH(A{i},\'TEST RESULTS\'!$B$5:$B$25,0))>=95,10,'
        f'IF(INDEX(\'TEST RESULTS\'!$C$5:$C$25,MATCH(A{i},\'TEST RESULTS\'!$B$5:$B$25,0))>=90,9,'
        f'IF(INDEX(\'TEST RESULTS\'!$C$5:$C$25,MATCH(A{i},\'TEST RESULTS\'!$B$5:$B$25,0))>=85,8,'
        f'IF(INDEX(\'TEST RESULTS\'!$C$5:$C$25,MATCH(A{i},\'TEST RESULTS\'!$B$5:$B$25,0))>=80,7,'
        f'IF(INDEX(\'TEST RESULTS\'!$C$5:$C$25,MATCH(A{i},\'TEST RESULTS\'!$B$5:$B$25,0))>=75,6,'
        f'IF(INDEX(\'TEST RESULTS\'!$C$5:$C$25,MATCH(A{i},\'TEST RESULTS\'!$B$5:$B$25,0))>=70,5,'
        f'IF(INDEX(\'TEST RESULTS\'!$C$5:$C$25,MATCH(A{i},\'TEST RESULTS\'!$B$5:$B$25,0))>=65,4,'
        f'IF(INDEX(\'TEST RESULTS\'!$C$5:$C$25,MATCH(A{i},\'TEST RESULTS\'!$B$5:$B$25,0))>=60,3,0))))))))),0)')
    
    calc.cell(i, 5, None)
    
    # FIX: Use INDEX/MATCH for the delta calculation as well
    calc.cell(i, 6,
        f'=IF(INDEX(\'TEST RESULTS\'!$C$5:$C$25,MATCH(A{i},\'TEST RESULTS\'!$B$5:$B$25,0))="",0,'
        f'IF(E{i}="",0,'
        f'IF(D{i}>E{i},0.1,'
        f'IF(D{i}<E{i},-0.1,0))))')
    
    calc.cell(i, 23, f'=A{i}&"  —  £"&TEXT(B{i},"0.0")&"m"')

calc["K1"] = "Manager"
calc["L1"] = "P1"
calc["M1"] = "P2"
calc["N1"] = "P3"
calc["O1"] = "P4"
calc["P1"] = "P5"
calc["Q1"] = "Captain"
calc["R1"] = "Team Value"
calc["S1"] = "Budget Left"
calc["T1"] = "GW Score"
calc["U1"] = "Status"

for r, name in enumerate(managers, start=2):
    calc.cell(r, 11, name)
    for c_off, ref in enumerate(["C10", "C11", "C12", "C13", "C14"]):
        calc.cell(r, 12 + c_off, f"='{name}'!{ref}")
    calc.cell(r, 17, f"='{name}'!C5")
    calc.cell(r, 18, f"='{name}'!C17")
    calc.cell(r, 19, f"='{name}'!C18")
    calc.cell(r, 20, f"='{name}'!C19")
    calc.cell(r, 21, f"='{name}'!C20")

# ---------------------------------------------------------
# HOME
# ---------------------------------------------------------
home = wb.create_sheet("HOME")
home.sheet_view.showGridLines = False
for col, w in {"A": 4, "B": 28, "C": 26, "D": 26, "E": 26, "F": 26, "G": 26, "H": 4}.items():
    home.column_dimensions[col].width = w

title(home, "YEAR 11 FANTASY LEAGUE", 1, 2, 7)

home.merge_cells("B2:G2")
home["B2"] = "GAMEWEEK 1"
home["B2"].font = Font(size=16, bold=True, color=BLUE)
home["B2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
home.row_dimensions[2].height = 32

home.merge_cells("B3:G3")
home.row_dimensions[3].height = 12

home.merge_cells("B4:G4")
home["B4"] = "DEADLINE  •  GAMEWEEK 1  •  TBC"
home["B4"].font = Font(size=13, bold=True, color=DARK)
home["B4"].fill = PatternFill("solid", fgColor=LIGHT_BLUE)
home["B4"].alignment = Alignment(horizontal="center", vertical="center")
home.row_dimensions[4].height = 34

home.merge_cells("B5:G5")
home["B5"] = "TEST DATE  •  GAMEWEEK 1  •  TBC"
home["B5"].font = Font(size=13, bold=True, color=DARK)
home["B5"].fill = PatternFill("solid", fgColor=LIGHT_BLUE)
home["B5"].alignment = Alignment(horizontal="center", vertical="center")
home.row_dimensions[5].height = 34

home.merge_cells("B6:G6")
home.row_dimensions[6].height = 12

home["B7"] = "CURRENT GAMEWEEK:"
home["B7"].font = Font(size=12, bold=True, color=GREY)
home["B7"].alignment = Alignment(horizontal="right", vertical="center")
home["C7"] = 1
home["C7"].font = Font(size=14, bold=True, color=NAVY)
home["C7"].fill = PatternFill("solid", fgColor=LIGHT_GOLD)
home["C7"].alignment = Alignment(horizontal="center", vertical="center")
home["C7"].border = Border(left=thin_grey, right=thin_grey, top=thin_grey, bottom=thin_grey)
home.row_dimensions[7].height = 30

home.merge_cells("B8:G8")
home.row_dimensions[8].height = 12

section(home, "HOW TO PLAY", 9, 2, 7)
steps = [
    "1.  Go to PICK YOUR TEAM and click your name.",
    "2.  Pick 5 players from the dropdowns (each shows the price).",
    "3.  Pick 1 captain using the GOLD CAPTAIN BOX at the top.",
    "4.  Stay under the £50m budget.",
    "5.  Enter your own test % on TEST RESULTS after each test.",
]
for i, text in enumerate(steps):
    r = 10 + i
    home.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
    c = home.cell(r, 2, text)
    box(c, PALE_BLUE)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    home.row_dimensions[r].height = 30

home.merge_cells("B16:G16")
home.row_dimensions[16].height = 12

section(home, "HOW POINTS WORK", 17, 2, 7)
points_explainer = [
    ("95–100%",   "10 points"),
    ("90–94%",    "9 points"),
    ("85–89%",    "8 points"),
    ("80–84%",    "7 points"),
    ("75–79%",    "6 points"),
    ("70–74%",    "5 points"),
    ("65–69%",    "4 points"),
    ("60–64%",    "3 points"),
    ("Below 60%", "0 points"),
]
for i, (label, val) in enumerate(points_explainer):
    r = 18 + i
    home.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    home.cell(r, 2, label).font = Font(bold=True, color=DARK, size=12)
    home.cell(r, 2).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    home.cell(r, 2).alignment = Alignment(horizontal="left", vertical="center", indent=1)
    home.cell(r, 2).border = Border(left=thin_grey, right=thin_grey, top=thin_grey, bottom=thin_grey)
    home.merge_cells(start_row=r, start_column=4, end_row=r, end_column=7)
    home.cell(r, 4, val).font = Font(color=DARK, size=12)
    home.cell(r, 4).alignment = Alignment(horizontal="left", vertical="center", indent=1)
    home.cell(r, 4).border = Border(left=thin_grey, right=thin_grey, top=thin_grey, bottom=thin_grey)
    home.row_dimensions[r].height = 26

home.merge_cells("B28:G28")
home.row_dimensions[28].height = 12

section(home, "HOW PRICES WORK", 29, 2, 7)
home.merge_cells("B30:G31")
home["B30"] = ("Prices move based on FORM — how each player performs compared to their\n"
               "OWN previous test. Consistent players don't move.")
home["B30"].font = Font(size=12, color=DARK)
home["B30"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
box(home["B30"], PALE_BLUE)
home.row_dimensions[30].height = 22
home.row_dimensions[31].height = 22

price_explainer = [
    ("Improved vs last test", "+£0.1m"),
    ("Same as last test",     "No change"),
    ("Declined vs last test", "-£0.1m"),
    ("First gameweek",        "No change"),
    ("Did not take test",     "No change"),
]
for i, (label, val) in enumerate(price_explainer):
    r = 32 + i
    home.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    home.cell(r, 2, label).font = Font(bold=True, color=DARK, size=12)
    home.cell(r, 2).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    home.cell(r, 2).alignment = Alignment(horizontal="left", vertical="center", indent=1)
    home.cell(r, 2).border = Border(left=thin_grey, right=thin_grey, top=thin_grey, bottom=thin_grey)
    home.merge_cells(start_row=r, start_column=5, end_row=r, end_column=7)
    home.cell(r, 5, val).font = Font(color=DARK, size=12)
    home.cell(r, 5).alignment = Alignment(horizontal="left", vertical="center", indent=1)
    home.cell(r, 5).border = Border(left=thin_grey, right=thin_grey, top=thin_grey, bottom=thin_grey)
    home.row_dimensions[r].height = 26

home.merge_cells("B38:G38")
home["B38"] = ("Prices stay between £7.0m and £13.0m. Small changes each week "
               "but add up over the season.")
home["B38"].font = Font(size=11, italic=True, color=GREY)
home["B38"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
home.row_dimensions[38].height = 30

# ---------------------------------------------------------
# PICK YOUR TEAM
# ---------------------------------------------------------
menu = wb.create_sheet("PICK YOUR TEAM")
menu.sheet_view.showGridLines = False
menu.column_dimensions["A"].width = 6
menu.column_dimensions["B"].width = 32
menu.column_dimensions["C"].width = 26

title(menu, "PICK YOUR TEAM", 1, 2, 3)

menu.merge_cells("B2:C2")
menu["B2"] = "Click your name to jump to your team sheet."
menu["B2"].font = Font(size=12, italic=True, color=GREY)
menu["B2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
menu.row_dimensions[2].height = 30

menu.merge_cells("B3:C3")
menu.row_dimensions[3].height = 12

menu["B4"] = "Manager"
menu["C4"] = "Status"
header_row(menu, 4, 2, 3)

for i, name in enumerate(managers):
    r = 5 + i
    c = menu.cell(r, 2, name)
    c.hyperlink = f"#'{name}'!A1"
    c.font = Font(bold=True, color=BLUE, underline="single", size=12)
    box(c, WHITE, True, BLUE)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    calc_row = 2 + i
    menu.cell(r, 3, f"=CALCULATIONS!U{calc_row}")
    box(menu.cell(r, 3), PALE_BLUE, align="center")
    menu.row_dimensions[r].height = 30

# ---------------------------------------------------------
# TEAM SHEETS
# ---------------------------------------------------------
def build_team_sheet(name):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 34
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 4
    ws.column_dimensions["F"].width = 4
    ws.column_dimensions["G"].width = 4

    title(ws, f"{name.upper()}'S TEAM", 1, 2, 4)

    ws.merge_cells("B2:D2")
    back = ws["B2"]
    back.value = "← Back to PICK YOUR TEAM"
    back.hyperlink = "#'PICK YOUR TEAM'!A1"
    back.font = Font(size=12, bold=True, color=BLUE, underline="single")
    back.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 30

    ws.row_dimensions[3].height = 12

    section(ws, "PICK YOUR CAPTAIN  (scores 2× points)", 4, 2, 4)
    ws["C5"] = ""
    box(ws["C5"], LIGHT_GOLD, True, DARK, "left")
    ws["C5"].font = Font(size=13, bold=True, color=NAVY)
    ws["C5"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[5].height = 34

    ws.merge_cells("B6:D6")
    ws["B6"] = "Step 1: Pick your 5 players below.  Step 2: Choose your captain in the GOLD box above."
    ws["B6"].font = Font(size=11, italic=True, color=GREY)
    ws["B6"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[6].height = 24

    ws.row_dimensions[7].height = 12

    section(ws, "SQUAD", 8, 2, 4)
    for c, h in enumerate(["Slot", "Player", "Price"], start=2):
        ws.cell(9, c, h)
    header_row(ws, 9, 2, 4)

    for r in range(10, 15):
        slot = r - 9
        ws.cell(r, 2, slot)
        ws.cell(r, 3, "")
        ws.cell(r, 4, f'=IF(C{r}="","",_xlfn.IFERROR(INDEX(\'CALCULATIONS\'!$C$2:$C$22,MATCH(C{r},\'CALCULATIONS\'!$W$2:$W$22,0)),""))')
        box(ws.cell(r, 2), PALE_BLUE, True, DARK, "center")
        box(ws.cell(r, 3), LIGHT_GOLD, False, DARK, "left")
        box(ws.cell(r, 4), WHITE, False, DARK, "center")
        ws.row_dimensions[r].height = 32

    ws.row_dimensions[15].height = 12

    section(ws, "TEAM SUMMARY", 16, 2, 4)
    summary = [
        ("Team Value",   "=SUM(D10:D14)"),
        ("Budget Left",  "=50-C17"),
        # FIX: Using SUMPRODUCT to correctly sum the array of 5 players, then adding captain
        ("GW Score",
         '=IF($C$20="VALID TEAM",'
         'SUMPRODUCT(SUMIF(\'CALCULATIONS\'!$W$2:$W$22,C10:C14,\'CALCULATIONS\'!$D$2:$D$22))'
         '+SUMIF(\'CALCULATIONS\'!$W$2:$W$22,C5,\'CALCULATIONS\'!$D$2:$D$22),0)'),
        ("Team Status",
         '=IF(COUNTBLANK(C10:C14)>0,"LESS THAN 5 PLAYERS",'
         'IF(SUM(D10:D14)>50,"OVER BUDGET",'
         'IF(C5="","NO CAPTAIN",'
         'IF(COUNTIF(C10:C14,C5)=0,"CAPTAIN NOT IN SQUAD",'
         'IF(SUMPRODUCT(COUNTIF(C10:C14,C10:C14))>5,"DUPLICATE PLAYER",'
         '"VALID TEAM")))))'),
    ]
    for i, (label, formula) in enumerate(summary):
        r = 17 + i
        ws.cell(r, 2, label)
        ws.cell(r, 3, formula)
        box(ws.cell(r, 2), LIGHT_BLUE, True, DARK, "left")
        box(ws.cell(r, 3), WHITE, True, DARK, "left")
        ws.row_dimensions[r].height = 30
    ws["C17"].number_format = '£0.0"m"'
    ws["C18"].number_format = '£0.0"m"'
    ws["C19"].number_format = '0'
    ws["C20"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.conditional_formatting.add("C20", FormulaRule(
        formula=['C20="VALID TEAM"'],
        fill=PatternFill("solid", fgColor=LIGHT_GREEN),
        font=Font(color=GREEN, bold=True)))
    ws.conditional_formatting.add("C20", FormulaRule(
        formula=['C20<>"VALID TEAM"'],
        fill=PatternFill("solid", fgColor=LIGHT_RED),
        font=Font(color=RED, bold=True)))

    ws.row_dimensions[21].height = 12

    ws.merge_cells("B22:D24")
    ws["B22"] = ("INPUT CELLS ARE GOLD.\n"
                 "Only the captain cell (C5) and the 5 player cells (C10:C14) should be edited.\n"
                 "Do NOT type 'CAPTAIN' anywhere in the squad — use the GOLD box above.")
    ws["B22"].font = Font(bold=True, color=DARK, size=12)
    ws["B22"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    box(ws["B22"], LIGHT_GOLD)

    ws.row_dimensions[25].height = 12

    section(ws, "EXAMPLE  (do not edit — for reference only)", 26, 2, 4, fill=SLATE)

    for c, h in enumerate(["Slot", "Player", "Price"], start=2):
        ws.cell(27, c, h)
    header_row(ws, 27, 2, 4, fill=SLATE_DARK)

    example = [
        (1, "Chinedu Chile  —  £11.0m",   "£11.0m"),
        (2, "Emily Warburton  —  £11.0m", "£11.0m"),
        (3, "Joe Paige  —  £10.5m",       "£10.5m"),
        (4, "Edwin Mayers  —  £9.0m",     "£9.0m"),
        (5, "Alfie Foley  —  £7.5m",      "£7.5m"),
    ]
    for i, (slot, p, price) in enumerate(example):
        r = 28 + i
        ws.cell(r, 2, slot)
        ws.cell(r, 3, p)
        ws.cell(r, 4, price)
        box(ws.cell(r, 2), PALE_BLUE, True, DARK, "center")
        box(ws.cell(r, 3), WHITE, False, DARK, "left")
        box(ws.cell(r, 4), WHITE, False, DARK, "center")
        ws.row_dimensions[r].height = 28

    ws.cell(33, 2, "CAPTAIN")
    box(ws.cell(33, 2), LIGHT_BLUE, True, DARK, "left")
    ws.merge_cells("C33:D33")
    ws.cell(33, 3, "Emily Warburton  —  £11.0m")
    box(ws.cell(33, 3), WHITE, True, DARK, "left")
    ws.row_dimensions[33].height = 28

    ws.cell(34, 2, "TEAM VALUE")
    box(ws.cell(34, 2), LIGHT_BLUE, True, DARK, "left")
    ws.merge_cells("C34:D34")
    ws.cell(34, 3, "£49.0m     BUDGET LEFT  £1.0m")
    box(ws.cell(34, 3), WHITE, True, DARK, "left")
    ws.row_dimensions[34].height = 28

    ws.cell(35, 2, "STATUS")
    box(ws.cell(35, 2), LIGHT_BLUE, True, DARK, "left")
    ws.merge_cells("C35:D35")
    ws.cell(35, 3, "✓ VALID TEAM")
    box(ws.cell(35, 3), LIGHT_GREEN, True, GREEN, "left")
    ws.row_dimensions[35].height = 28

    dv_players = DataValidation(
        type="list",
        formula1="=CALCULATIONS!$W$2:$W$22",
        allow_blank=True,
    )
    dv_players.error = "Please choose a player from the dropdown list."
    dv_players.errorTitle = "Invalid player"
    dv_players.prompt = "Click the arrow to choose a player."
    dv_players.promptTitle = "Pick a player"
    dv_players.showErrorMessage = True
    dv_players.showInputMessage = True
    ws.add_data_validation(dv_players)
    dv_players.add("C10:C14")

    dv_captain = DataValidation(
        type="list",
        formula1=f"='{name}'!$C$10:$C$14",
        allow_blank=True,
    )
    dv_captain.error = "Choose one of your 5 players."
    dv_captain.errorTitle = "Invalid captain"
    dv_captain.prompt = "Pick your captain from your 5 players."
    dv_captain.promptTitle = "Pick a captain"
    dv_captain.showErrorMessage = True
    dv_captain.showInputMessage = True
    ws.add_data_validation(dv_captain)
    dv_captain.add("C5")

    return ws

for name in managers:
    build_team_sheet(name)

# ---------------------------------------------------------
# TEST RESULTS
# ---------------------------------------------------------
results = wb.create_sheet("TEST RESULTS")
results.sheet_view.showGridLines = False
results.column_dimensions["A"].width = 4
results.column_dimensions["B"].width = 32
results.column_dimensions["C"].width = 42

title(results, "TEST RESULTS", 1, 2, 3)

results.merge_cells("B2:C2")
results["B2"] = "GAMEWEEK 1  •  Test on TBC"
results["B2"].font = Font(size=12, bold=True, color=BLUE)
results["B2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
results.row_dimensions[2].height = 30

results.merge_cells("B3:C3")
results.row_dimensions[3].height = 12

results["B4"] = "Student"
results["C4"] = "Test %  —  type 90, not 90%"
header_row(results, 4, 2, 3)

for i, name in enumerate(managers):
    r = 5 + i
    results.cell(r, 2, name)
    results.cell(r, 3, "")
    box(results.cell(r, 2), WHITE, False, DARK, "left")
    box(results.cell(r, 3), LIGHT_GOLD, False, DARK, "center")
    results.cell(r, 3).number_format = '0"%"'
    results.row_dimensions[r].height = 30

dv_pct = DataValidation(
    type="whole", operator="between",
    formula1="0", formula2="100",
    allow_blank=True,
)
dv_pct.error = "Enter a whole number from 0 to 100."
dv_pct.errorTitle = "Invalid percentage"
dv_pct.showErrorMessage = True
results.add_data_validation(dv_pct)
dv_pct.add("C5:C25")

results.merge_cells("B28:C30")
results["B28"] = ("Enter a whole number between 0 and 100.\n"
                  "Do NOT type the % sign — just the number. Example: 90")
results["B28"].font = Font(bold=True, color=DARK, size=12)
results["B28"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
box(results["B28"], LIGHT_GOLD)

# ---------------------------------------------------------
# PLAYER MARKET
# ---------------------------------------------------------
market = wb.create_sheet("PLAYER MARKET")
market.sheet_view.showGridLines = False
market.column_dimensions["A"].width = 4
market.column_dimensions["B"].width = 34
market.column_dimensions["C"].width = 14
market.column_dimensions["D"].width = 14
market.column_dimensions["E"].width = 20

title(market, "PLAYER MARKET", 1, 2, 5)

market.merge_cells("B2:E2")
market["B2"] = "Prices  •  This week's change  •  Ownership"
market["B2"].font = Font(size=12, color=GREY, italic=True)
market["B2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
market.row_dimensions[2].height = 30

market.merge_cells("B3:E3")
market.row_dimensions[3].height = 12

for c, h in enumerate(["Player", "Price", "Change", "Ownership %"], start=2):
    market.cell(4, c, h)
header_row(market, 4, 2, 5)

for i, idx in enumerate(range(2, 23)):
    r = 5 + i
    market.cell(r, 2, f"=CALCULATIONS!A{idx}")
    market.cell(r, 3, f"=CALCULATIONS!C{idx}")
    market.cell(r, 4, f"=CALCULATIONS!F{idx}")
    market.cell(r, 5,
        f"=COUNTIF('CALCULATIONS'!$L$2:$P$22,CALCULATIONS!W{idx})"
        f"/COUNTA('CALCULATIONS'!$K$2:$K$22)")
    box(market.cell(r, 2), WHITE if i % 2 == 0 else PALE_BLUE, False, DARK, "left")
    box(market.cell(r, 3), WHITE if i % 2 == 0 else PALE_BLUE, True, DARK, "center")
    box(market.cell(r, 4), WHITE if i % 2 == 0 else PALE_BLUE, True, DARK, "center")
    box(market.cell(r, 5), WHITE if i % 2 == 0 else PALE_BLUE, False, DARK, "center")
    market.cell(r, 3).number_format = '£0.0"m"'
    market.cell(r, 4).number_format = '"▲ +£"0.0"m";"▼ -£"0.0"m";"— £0.0m"'
    market.cell(r, 5).number_format = '0%'
    market.row_dimensions[r].height = 28

market.conditional_formatting.add("D5:D25", FormulaRule(
    formula=['D5>0'],
    fill=PatternFill("solid", fgColor=LIGHT_GREEN),
    font=Font(color=GREEN, bold=True)))
market.conditional_formatting.add("D5:D25", FormulaRule(
    formula=['D5<0'],
    fill=PatternFill("solid", fgColor=LIGHT_RED),
    font=Font(color=RED, bold=True)))
market.conditional_formatting.add("D5:D25", FormulaRule(
    formula=['D5=0'],
    fill=PatternFill("solid", fgColor="F3F4F6"),
    font=Font(color=GREY, bold=True)))

market.merge_cells("B28:E30")
market["B28"] = ("PRICE RULES\n"
                 "Prices move ±£0.1m based on form vs each player's previous test.\n"
                 "Minimum £7.0m  •  Maximum £13.0m")
market["B28"].font = Font(bold=True, color=DARK, size=12)
market["B28"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
box(market["B28"], LIGHT_BLUE)

# ---------------------------------------------------------
# 1. RAW RESULTS (hidden)
# ---------------------------------------------------------
raw = wb.create_sheet("1. Raw Results")
raw.sheet_state = "hidden"
raw.column_dimensions["A"].width = 4
raw.column_dimensions["B"].width = 14
raw.column_dimensions["C"].width = 32
raw.column_dimensions["D"].width = 14

raw["B1"] = "Gameweek"
raw["C1"] = "Manager"
raw["D1"] = "Points"
header_row(raw, 1, 2, 4)

for i, name in enumerate(managers):
    r = 2 + i
    raw.cell(r, 2, "GW1")
    raw.cell(r, 3, name)
    raw.cell(r, 4, f"=CALCULATIONS!T{2 + i}")
    box(raw.cell(r, 2), WHITE, True, DARK, "center")
    box(raw.cell(r, 3), WHITE, False, DARK, "left")
    box(raw.cell(r, 4), WHITE, True, DARK, "center")
    raw.row_dimensions[r].height = 28

# ---------------------------------------------------------
# 2. OVERALL LEADERBOARD
# ---------------------------------------------------------
overall = wb.create_sheet("2. Overall Leaderboard")
overall.sheet_view.showGridLines = False
overall.column_dimensions["A"].width = 4
overall.column_dimensions["B"].width = 12
overall.column_dimensions["C"].width = 32
overall.column_dimensions["D"].width = 18

title(overall, "OVERALL LEADERBOARD", 1, 2, 4)

overall.merge_cells("B2:D2")
overall["B2"] = "All gameweeks combined  •  Sorted by points (shared rank on ties)"
overall["B2"].font = Font(size=12, color=GREY, italic=True)
overall["B2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
overall.row_dimensions[2].height = 26

overall.merge_cells("B3:D3")
overall["B3"] = (
    '=IF(MAX($D$6:$D$26)=0,"🏆  OVERALL LEADER:  No results yet",'
    '"🏆  OVERALL LEADER:  "&INDEX($C$6:$C$26,MATCH(MAX($D$6:$D$26),$D$6:$D$26,0))'
    '&"  —  "&MAX($D$6:$D$26)&" pts")'
)
overall["B3"].font = Font(size=14, bold=True, color="92400E")
overall["B3"].fill = PatternFill("solid", fgColor=LIGHT_GOLD)
overall["B3"].alignment = Alignment(horizontal="center", vertical="center")
overall["B3"].border = Border(left=thin_grey, right=thin_grey, top=thin_grey, bottom=thin_grey)
overall.row_dimensions[3].height = 38

overall.merge_cells("B4:D4")
overall.row_dimensions[4].height = 12

for c, h in enumerate(["Rank", "Manager", "Overall Points"], start=2):
    overall.cell(5, c, h)
header_row(overall, 5, 2, 4)

for i in range(21):
    r = 6 + i
    fill = WHITE if i % 2 == 0 else PALE_BLUE
    box(overall.cell(r, 2), fill, True, DARK, "center")
    box(overall.cell(r, 3), fill, False, DARK, "left")
    box(overall.cell(r, 4), fill, True, DARK, "center")
    overall.row_dimensions[r].height = 30

overall.conditional_formatting.add("B6:B26", FormulaRule(
    formula=['$B6=1'],
    fill=PatternFill("solid", fgColor=LIGHT_GOLD),
    font=Font(bold=True, color="92400E")))
overall.conditional_formatting.add("B6:B26", FormulaRule(
    formula=['$B6=2'],
    fill=PatternFill("solid", fgColor="E5E7EB"),
    font=Font(bold=True, color="374151")))
overall.conditional_formatting.add("B6:B26", FormulaRule(
    formula=['$B6=3'],
    fill=PatternFill("solid", fgColor="FED7AA"),
    font=Font(bold=True, color="9A3412")))

# ---------------------------------------------------------
# 3. GW LEADERBOARD
# ---------------------------------------------------------
gw = wb.create_sheet("3. GW Leaderboard")
gw.sheet_view.showGridLines = False
gw.column_dimensions["A"].width = 4
gw.column_dimensions["B"].width = 12
gw.column_dimensions["C"].width = 32
gw.column_dimensions["D"].width = 18

title(gw, "GAMEWEEK LEADERBOARD", 1, 2, 4)

gw.merge_cells("B2:D2")
gw["B2"] = "This gameweek only  •  Sorted by points (shared rank on ties)"
gw["B2"].font = Font(size=12, color=GREY, italic=True)
gw["B2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
gw.row_dimensions[2].height = 26

gw["B3"] = "CURRENT GAMEWEEK:"
gw["B3"].font = Font(size=12, bold=True, color=GREY)
gw["B3"].alignment = Alignment(horizontal="right", vertical="center")
gw["C3"] = 1
gw["C3"].font = Font(size=14, bold=True, color=NAVY)
gw["C3"].fill = PatternFill("solid", fgColor=LIGHT_GOLD)
gw["C3"].alignment = Alignment(horizontal="center", vertical="center")
gw["C3"].border = Border(left=thin_grey, right=thin_grey, top=thin_grey, bottom=thin_grey)
gw.row_dimensions[3].height = 30

gw.merge_cells("B4:D4")
gw.row_dimensions[4].height = 10

gw.merge_cells("B5:D5")
gw["B5"] = (
    '=IF(MAX($D$8:$D$28)=0,"🏆  GW"&$C$3&" WINNER:  No results yet",'
    '"🏆  GW"&$C$3&" WINNER:  "&INDEX($C$8:$C$28,MATCH(MAX($D$8:$D$28),$D$8:$D$28,0))'
    '&"  —  "&MAX($D$8:$D$28)&" pts")'
)
gw["B5"].font = Font(size=14, bold=True, color="92400E")
gw["B5"].fill = PatternFill("solid", fgColor=LIGHT_GOLD)
gw["B5"].alignment = Alignment(horizontal="center", vertical="center")
gw["B5"].border = Border(left=thin_grey, right=thin_grey, top=thin_grey, bottom=thin_grey)
gw.row_dimensions[5].height = 38

gw.merge_cells("B6:D6")
gw.row_dimensions[6].height = 12

for c, h in enumerate(["Rank", "Manager", "GW Points"], start=2):
    gw.cell(7, c, h)
header_row(gw, 7, 2, 4)

for i in range(21):
    r = 8 + i
    fill = WHITE if i % 2 == 0 else PALE_BLUE
    box(gw.cell(r, 2), fill, True, DARK, "center")
    box(gw.cell(r, 3), fill, False, DARK, "left")
    box(gw.cell(r, 4), fill, True, DARK, "center")
    gw.row_dimensions[r].height = 30

gw.conditional_formatting.add("B8:B28", FormulaRule(
    formula=['$B8=1'],
    fill=PatternFill("solid", fgColor=LIGHT_GOLD),
    font=Font(bold=True, color="92400E")))
gw.conditional_formatting.add("B8:B28", FormulaRule(
    formula=['$B8=2'],
    fill=PatternFill("solid", fgColor="E5E7EB"),
    font=Font(bold=True, color="374151")))
gw.conditional_formatting.add("B8:B28", FormulaRule(
    formula=['$B8=3'],
    fill=PatternFill("solid", fgColor="FED7AA"),
    font=Font(bold=True, color="9A3412")))

# ---------------------------------------------------------
# RULES
# ---------------------------------------------------------
rules = wb.create_sheet("RULES")
rules.sheet_view.showGridLines = False
rules.column_dimensions["A"].width = 4
rules.column_dimensions["B"].width = 28
rules.column_dimensions["C"].width = 72

title(rules, "RULES & HOW TO PLAY", 1, 2, 3)

rule_rows = [
    ("BUDGET",
     "£50.0m per manager. Pick exactly 5 players. Squad value cannot exceed £50.0m."),
    ("CAPTAIN",
     "Choose 1 captain using the GOLD dropdown at the top of your team sheet. The captain scores 2× their normal Gameweek points."),
    ("DEADLINE",
     "Set manually on HOME each gameweek."),
    ("TEST DATE",
     "Set manually on HOME each gameweek."),
    ("RESULT ENTRY",
     "After the test, each student enters only their own percentage on TEST RESULTS."),
    ("POINTS",
     "95–100% = 10 • 90–94% = 9 • 85–89% = 8 • 80–84% = 7 • 75–79% = 6 • 70–74% = 5 • 65–69% = 4 • 60–64% = 3 • Below 60% = 0."),
    ("PRICE CHANGES",
     "Prices move based on FORM. If a student improves vs their last test → +£0.1m. If they decline → -£0.1m. Same → no change. First week or blank test → no change."),
    ("PRICE LIMITS",
     "Minimum £7.0m. Maximum £13.0m."),
    ("TRANSFERS",
     "After every Gameweek: unlimited transfers, free transfers, no penalties, captains can change."),
    ("TRANSFER DEADLINE",
     "The transfer window closes at the next Gameweek deadline."),
    ("GOLDEN RULE",
     "Users only need to pick 5 players, pick a captain and enter their own test result. Excel handles the rest."),
]

for i, (topic, explanation) in enumerate(rule_rows):
    r = 4 + (i * 2)
    rules.cell(r, 2, topic)
    rules.cell(r, 3, explanation)
    box(rules.cell(r, 2), LIGHT_BLUE, True, DARK, "left")
    box(rules.cell(r, 3), WHITE, False, DARK, "left")
    rules.cell(r, 3).alignment = Alignment(
        horizontal="left", vertical="center", wrap_text=True, indent=1)
    rules.row_dimensions[r].height = 40
    rules.row_dimensions[r + 1].height = 8

# ---------------------------------------------------------
# TAB COLOURS
# ---------------------------------------------------------
tab_colours = {
    "HOME": BLUE,
    "PICK YOUR TEAM": "1D4ED8",
    "TEST RESULTS": GOLD,
    "PLAYER MARKET": "2563EB",
    "2. Overall Leaderboard": "0F766E",
    "3. GW Leaderboard": "0D9488",
    "RULES": SLATE,
    "CALCULATIONS": "64748B",
    "1. Raw Results": "94A3B8",
}
for ws in wb.worksheets:
    if ws.title in tab_colours:
        ws.sheet_properties.tabColor = tab_colours[ws.title]

for name in managers:
    wb[name].sheet_properties.tabColor = SLATE

for ws in wb.worksheets:
    ws.sheet_view.zoomScale = 90

main_order = [
    "HOME",
    "PICK YOUR TEAM",
    "TEST RESULTS",
    "PLAYER MARKET",
    "2. Overall Leaderboard",
    "3. GW Leaderboard",
    "RULES",
    "1. Raw Results",
    "CALCULATIONS",
]
ordered = main_order + managers
wb._sheets = [wb[t] for t in ordered]

wb.active = wb.index(wb["HOME"])

# ---------------------------------------------------------
# SAVE
# ---------------------------------------------------------
path = r"C:\Users\chine\OneDrive\Python Stuff\Fantasy League\Year_11_Fantasy_League.xlsx"
wb.save(path)
print(f"Built: {path}")
print(f"Managers: {len(managers)}")
print(f"Players: {len(players)}")
print()
print("=" * 60)
print("NEXT STEPS:")
print("1. Open the workbook in Excel. Accept the repair.")
print("2. Paste formulas into: 2. Overall Leaderboard!B3 and B6")
print("3. Paste formulas into: 3. GW Leaderboard!B5 and B8")
print("4. Save and close.")
print("=" * 60)
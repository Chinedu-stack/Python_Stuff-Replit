from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill, Font, Border, Side, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

# =========================================================
# CONFIG
# =========================================================
FILE = r"C:\Users\chine\OneDrive\Python Stuff\Fantasy League\Year_11_Fantasy_League.xlsx"

WHITE = "FFFFFF"
PALE_BLUE = "F5F9FF"
DARK = "172033"
LIGHT_GREY = "E5E7EB"
thin_grey = Side(style="thin", color=LIGHT_GREY)

def box_cell(cell, fill=WHITE, bold=False, color=DARK, align="center"):
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(bold=bold, color=color, size=12)
    cell.alignment = Alignment(horizontal=align, vertical="center")
    cell.border = Border(left=thin_grey, right=thin_grey, top=thin_grey, bottom=thin_grey)

# =========================================================
# LOAD
# =========================================================
wb_values = load_workbook(FILE, data_only=True)
wb = load_workbook(FILE)

home = wb["HOME"]
current_gw = home["C7"].value

if current_gw is None:
    print("ERROR: HOME!C7 is empty. Can't determine current gameweek.")
    raise SystemExit(1)

next_gw = current_gw + 1

print(f"Loaded: {FILE}")
print(f"Current gameweek: {current_gw}")
print(f"Advancing to: {next_gw}")
print()

# ---------------------------------------------------------
# GUARD: cached values present?
# ---------------------------------------------------------
probe = wb_values["CALCULATIONS"]["D2"].value
if probe is None:
    print("WARNING: No cached values found in the workbook.")
    print("         Open the file in Excel, save it (Ctrl+S), close it,")
    print("         then re-run this script.")
    print("         Aborting so we don't corrupt history.")
    raise SystemExit(1)

# =========================================================
# 0. MANAGER LIST
# =========================================================
menu = wb["PICK YOUR TEAM"]
manager_names = []
for r in range(5, 26):
    v = menu.cell(r, 2).value
    if v:
        manager_names.append(v)

print(f"Managers detected: {len(manager_names)}")

# =========================================================
# 1. FREEZE TEST RESULTS
# =========================================================
results = wb["TEST RESULTS"]
results_v = wb_values["TEST RESULTS"]

freeze_col = 3 + current_gw
freeze_letter = get_column_letter(freeze_col)

src_cell = results.cell(4, 3)
src_font = src_cell.font.copy()
src_fill = src_cell.fill.copy()
src_align = src_cell.alignment.copy()
src_border = src_cell.border.copy()

for r in range(5, 26):
    results.cell(r, freeze_col).value = results_v.cell(r, 3).value

hdr = results.cell(4, freeze_col)
hdr.value = f"GW{current_gw} %"
hdr.font = src_font
hdr.fill = src_fill
hdr.alignment = src_align
hdr.border = src_border

results.cell(4, 3).value = f"Test %  —  GW{next_gw}  —  type 90, not 90%"

for r in range(5, 26):
    results.cell(r, 3).value = None

print(f"  TEST RESULTS: GW{current_gw} frozen in column {freeze_letter}, column C cleared.")

# =========================================================
# 2. FREEZE POINTS THIS WEEK INTO "LAST WEEK PTS"
# =========================================================
calc = wb["CALCULATIONS"]
calc_v = wb_values["CALCULATIONS"]

# E (Last Week Pts) <- D (this week's points, from cached file)
for r in range(2, 23):
    calc.cell(r, 5).value = calc_v.cell(r, 4).value

print(f"  CALCULATIONS: Last Week Pts updated from this week's points.")

# =========================================================
# 3. FREEZE Δ INTO GROWING COLUMNS
# =========================================================
# Layout: G=GW1 Δ, H=GW2 Δ, I=GW3 Δ, ...
freeze_offset = current_gw - 1
delta_col = 7 + freeze_offset      # G=7, H=8, I=9...

# Copy header format from F1 (Δ)
delta_hdr_src = calc.cell(1, 6)
delta_hdr_font = delta_hdr_src.font.copy()
delta_hdr_fill = delta_hdr_src.fill.copy()
delta_hdr_align = delta_hdr_src.alignment.copy()
delta_hdr_border = delta_hdr_src.border.copy()

for r in range(2, 23):
    calc.cell(r, delta_col).value = calc_v.cell(r, 6).value

dh = calc.cell(1, delta_col)
dh.value = f"GW{current_gw} Δ"
dh.font = delta_hdr_font
dh.fill = delta_hdr_fill
dh.alignment = delta_hdr_align
dh.border = delta_hdr_border

print(f"  CALCULATIONS: GW{current_gw} Δ frozen in {get_column_letter(delta_col)}.")

# =========================================================
# 4. REBUILD CURRENT PRICE FORMULA
# =========================================================
# Current Price = MIN(13, MAX(7, Base + sum of all frozen Δs))
delta_cols = []
for gw in range(1, current_gw + 1):
    off = gw - 1
    delta_cols.append(get_column_letter(7 + off))   # G, H, I...

for r in range(2, 23):
    if delta_cols:
        delta_sum = "+".join(f"{c}{r}" for c in delta_cols)
        calc.cell(r, 3).value = f"=MIN(13,MAX(7,B{r}+{delta_sum}))"
    else:
        calc.cell(r, 3).value = f"=MIN(13,MAX(7,B{r}))"

print(f"  CALCULATIONS: Current Price formula rebuilt (Base + {'+'.join(delta_cols) if delta_cols else 'nothing'}).")

# =========================================================
# 5. APPEND GW SCORES TO "1. Raw Results"
# =========================================================
raw = wb["1. Raw Results"]
raw_v = wb_values["1. Raw Results"]

first_empty = 2
while raw.cell(first_empty, 3).value not in (None, ""):
    first_empty += 1

src_gw = raw.cell(2, 2)
src_mgr = raw.cell(2, 3)
src_pts = raw.cell(2, 4)

for i, name in enumerate(manager_names):
    r = first_empty + i
    calc_row = 2 + i

    gw_cell = raw.cell(r, 2)
    gw_cell.value = f"GW{current_gw}"
    gw_cell.font = src_gw.font.copy()
    gw_cell.fill = src_gw.fill.copy()
    gw_cell.alignment = src_gw.alignment.copy()
    gw_cell.border = src_gw.border.copy()

    mgr_cell = raw.cell(r, 3)
    mgr_cell.value = name
    mgr_cell.font = src_mgr.font.copy()
    mgr_cell.fill = src_mgr.fill.copy()
    mgr_cell.alignment = src_mgr.alignment.copy()
    mgr_cell.border = src_mgr.border.copy()

    pts_cell = raw.cell(r, 4)
    cached = raw_v.cell(r, 4).value
    if cached is None:
        cached = calc_v.cell(calc_row, 20).value
    pts_cell.value = cached if cached is not None else 0
    pts_cell.font = src_pts.font.copy()
    pts_cell.fill = src_pts.fill.copy()
    pts_cell.alignment = src_pts.alignment.copy()
    pts_cell.border = src_pts.border.copy()

    raw.row_dimensions[r].height = 28

print(f"  1. RAW RESULTS: appended {len(manager_names)} rows for GW{current_gw} "
      f"(rows {first_empty}-{first_empty + len(manager_names) - 1}).")

# =========================================================
# 6. UPDATE GW LEADERBOARD SELECTOR
# =========================================================
gw_lb = wb["3. GW Leaderboard"]
gw_lb["C3"].value = next_gw
print(f"  3. GW LEADERBOARD: selector updated to GW{next_gw}.")

# =========================================================
# 7. CLEAR TEAM INPUTS + RE-APPLY DROPDOWNS (PLAYER + CAPTAIN)
# =========================================================
for name in manager_names:
    if name not in wb.sheetnames:
        continue
    ws = wb[name]
    
    # Clear current inputs
    ws["C5"].value = None
    for r in range(10, 15):
        ws.cell(r, 3).value = None

    # Remove existing dropdowns for C5 and C10:C14 to avoid duplicates
    ws.data_validations.dataValidation = [
        dv for dv in ws.data_validations.dataValidation
        if str(dv.sqref) not in ("C5", "C10:C14")
    ]

    # Re-apply PLAYER dropdown (C10:C14)
    dv_players = DataValidation(
        type="list",
        formula1="=CALCULATIONS!$W$2:$W$22",
        allow_blank=True,
    )
    dv_players.error = "Please choose a player from the dropdown list."
    dv_players.errorTitle = "Invalid player"
    dv_players.prompt = "Click the arrow to choose a player."
    dv_players.promptTitle = "Pick a player"
    dv_players.showErrorMessage = True
    dv_players.showInputMessage = True
    ws.add_data_validation(dv_players)
    dv_players.add("C10:C14")

    # Re-apply CAPTAIN dropdown (C5) - only shows the 5 picked players
    dv_captain = DataValidation(
        type="list",
        formula1=f"='{name}'!$C$10:$C$14",
        allow_blank=True,
    )
    dv_captain.error = "Choose one of your 5 players."
    dv_captain.errorTitle = "Invalid captain"
    dv_captain.prompt = "Pick your captain from your 5 players."
    dv_captain.promptTitle = "Pick a captain"
    dv_captain.showErrorMessage = True
    dv_captain.showInputMessage = True
    ws.add_data_validation(dv_captain)
    dv_captain.add("C5")

print(f"  Cleared {len(manager_names)} team sheets and re-applied player + captain dropdowns.")

# =========================================================
# 8. UPDATE HEADERS — PLACEHOLDERS
# =========================================================
home["C7"].value = next_gw
home["B2"].value = f"GAMEWEEK {next_gw}"
home["B4"].value = f"DEADLINE  •  GAMEWEEK {next_gw}  •  TBC"
home["B5"].value = f"TEST DATE  •  GAMEWEEK {next_gw}  •  TBC"

results["B2"].value = f"GAMEWEEK {next_gw}  •  Test on TBC"

print(f"  HOME + TEST RESULTS: headers set to GW{next_gw} with TBC placeholders.")
print(f"  → Remember to edit the deadline, test date, and subject manually.")

# =========================================================
# 9. SAVE
# =========================================================
wb.active = wb.index(wb["HOME"])
wb.save(FILE)

print()
print(f"Saved. File is now on Gameweek {next_gw}.")
print()
print("REMINDER: Open the file in Excel and save it once so the next")
print("          advance has fresh cached values to work with.")